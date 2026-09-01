from dataclasses import replace
from datetime import datetime, timedelta, timezone

import pytest

from app.services.report_export import DatasetReportExporter
from app.services import DataAnalysisOrchestrator
from app.adapters import build_mock_adapters
from app.config import Settings
from app.domain.models import (
    AgentResponse,
    CanonicalAnalysisRequest,
    PrimaryIntent,
    TrustedIdentity,
)
from app.intent import RuleBasedIntentClassifier
from app.services.dataset_followup import DatasetLifecycleCleaner
from app.stores import InMemorySessionStore
from minio_followup_store import DatasetReference, DatasetScope, LoadedDataset


class Store:
    def load_dataset(self, reference, *, current_scope):
        assert current_scope == reference.scope
        return LoadedDataset(
            reference,
            ({"月份": "2026-07", "销售额": 100}, {"月份": "2026-08", "销售额": 120}),
        )


class Minio:
    def __init__(self):
        self.objects = {}

    def put_object(self, bucket, name, data, length, **kwargs):
        payload = data.read()
        assert len(payload) == length
        self.objects[(bucket, name)] = payload
        self.metadata = kwargs.get("metadata")

    def presigned_get_object(self, bucket, name, expires):
        return f"http://minio/{bucket}/{name}"

    def remove_object(self, bucket, name):
        self.objects.pop((bucket, name), None)


def reference():
    now_dt = datetime.now(timezone.utc)
    now = now_dt.isoformat()
    scope = DatasetScope("tenant", "user", "app", "conversation")
    return DatasetReference(
        dataset_id="dataset-1", bucket="bam", object_name="dataset",
        scope=scope, columns=("月份", "销售额"), row_count=2, byte_size=10,
        snapshot_id="snapshot", data_as_of=now, created_at=now,
        expires_at=(now_dt + timedelta(hours=1)).isoformat(),
        source_type="TEST", source_ref="test",
    )


@pytest.mark.parametrize(
    ("file_format", "signature"),
    [("xlsx", b"PK"), ("docx", b"PK"), ("pdf", b"%PDF")],
)
def test_exports_real_report_file(file_format, signature):
    minio = Minio()
    ref = reference()
    result = DatasetReportExporter(minio, Store(), bucket="bam").export(
        ref, scope=ref.scope, file_format=file_format, title="销售分析报告"
    )
    payload = minio.objects[("bam", result["object_name"])]
    assert payload.startswith(signature)
    assert result["download_url"].startswith("http://minio/")
    assert result["report_id"].startswith("report-")
    assert result["report_reference"]["content_sha256"]
    assert minio.metadata["expires-at"] == result["object_expires_at"]


def test_xlsx_export_neutralizes_formula_injection_and_illegal_controls():
    from io import BytesIO
    from openpyxl import load_workbook

    exporter = DatasetReportExporter(Minio(), Store(), bucket="bam")
    payload = exporter._xlsx(
        ["value"],
        [{"value": "=HYPERLINK(\"https://example.invalid\")\x00"}],
        "+unsafe title",
    )
    workbook = load_workbook(BytesIO(payload), data_only=False)
    try:
        assert workbook["数据"]["A2"].data_type == "s"
        assert workbook["数据"]["A2"].value.startswith("'=")
        assert "\x00" not in workbook["数据"]["A2"].value
        assert workbook["说明"]["B1"].value == "'+unsafe title"
    finally:
        workbook.close()


def test_composite_xlsx_keeps_independent_datasets_on_separate_sheets():
    from io import BytesIO
    from openpyxl import load_workbook

    minio = Minio()
    first = reference()
    second = replace(
        first,
        dataset_id="dataset-2",
        object_name="dataset-2",
        scope=DatasetScope("tenant", "user", "app", "dag-task-2"),
        snapshot_id="snapshot-2",
    )
    exporter = DatasetReportExporter(minio, Store(), bucket="bam")
    result = exporter.export_many(
        [("销售趋势", first), ("医院覆盖", second)],
        scope=DatasetScope("tenant", "user", "app", "root"),
        file_format="xlsx",
        title="综合分析报告",
    )

    payload = minio.objects[("bam", result["object_name"])]
    workbook = load_workbook(BytesIO(payload), data_only=True, read_only=True)
    try:
        assert workbook.sheetnames == ["报告说明", "章节1", "章节2"]
        assert workbook["报告说明"]["B3"].value == 2
        assert workbook["章节1"]["B2"].value == "dataset-1"
        assert workbook["章节2"]["B2"].value == "dataset-2"
    finally:
        workbook.close()
    assert result["dataset_ids"] == ["dataset-1", "dataset-2"]


@pytest.mark.asyncio
async def test_expired_report_object_and_redis_reference_are_cleaned():
    minio = Minio()
    minio.objects[("bam", "data-analysis/reports/expired.xlsx")] = b"expired"
    sessions = InMemorySessionStore()
    expires_at = (datetime.now(timezone.utc) - timedelta(seconds=1)).isoformat()
    await sessions.put_report_reference({
        "report_id": "report-expired",
        "bucket": "bam",
        "object_name": "data-analysis/reports/expired.xlsx",
        "expires_at": expires_at,
    })
    cleaner = DatasetLifecycleCleaner(
        Store(), sessions, interval_seconds=300, batch_size=10,
        report_client=minio, report_bucket="bam",
    )
    assert await cleaner.cleanup_once() == 1
    assert ("bam", "data-analysis/reports/expired.xlsx") not in minio.objects
    assert await sessions.list_expired_report_references(
        now_epoch=datetime.now(timezone.utc).timestamp(), limit=10
    ) == []


@pytest.mark.asyncio
async def test_report_is_created_inside_chat_workflow_on_explicit_request():
    minio = Minio()
    sessions = InMemorySessionStore()
    ref = reference()
    await sessions.put_dataset_reference(ref.to_dict(), recent_limit=10)
    agent = DataAnalysisOrchestrator(
        settings=Settings(env="test", adapter_mode="mock", intent_model_enabled=False),
        classifier=RuleBasedIntentClassifier(),
        adapters=build_mock_adapters(),
        sessions=sessions,
        report_exporter=DatasetReportExporter(minio, Store(), bucket="bam"),
    )
    request = CanonicalAnalysisRequest(
        conversation_id="conversation", application_id="app",
        tenant_id="tenant", user_id="user", original_question="生成PDF分析报告",
        primary_intent=PrimaryIntent.REPORT_GENERATION,
    )
    response = AgentResponse(
        request_id=request.request_id,
        conversation_id=request.conversation_id,
        status="COMPLETED",
        intent=request.primary_intent,
        answer="分析完成。",
    )

    await agent._attach_requested_report(
        response,
        request=request,
        identity=TrustedIdentity(tenant_id="tenant", user_id="user"),
        dataset_id=ref.dataset_id,
    )

    assert response.files[0].format == "pdf"
    assert response.files[0].download_url.startswith("http://minio/")
    assert response.files[0].dataset_id == ref.dataset_id
