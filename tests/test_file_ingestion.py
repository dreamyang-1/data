import io
import zipfile
from datetime import datetime, timezone

import pytest
from openpyxl import Workbook

from app.services.file_ingestion import FileImportError, SpreadsheetFileImporter
from app.services.dataset_followup import (
    is_dataset_operation_followup,
    plan_dataset_followup,
)
from minio_followup_store import DatasetReference, DatasetScope


class Response(io.BytesIO):
    def close(self):
        super().close()

    def release_conn(self):
        pass


class Minio:
    def __init__(self, payload):
        self.payload = payload

    def stat_object(self, _bucket, _name):
        return type("Stat", (), {"size": len(self.payload)})()

    def get_object(self, _bucket, _name):
        return Response(self.payload)


class Store:
    def save_dataset(self, **kwargs):
        now = datetime.now(timezone.utc).isoformat()
        return DatasetReference(
            dataset_id="dataset-upload-1",
            bucket="bam",
            object_name="datasets/1",
            scope=kwargs["scope"],
            columns=tuple(kwargs["columns"]),
            row_count=len(kwargs["rows"]),
            byte_size=100,
            snapshot_id=kwargs["snapshot_id"],
            data_as_of=now,
            created_at=now,
            expires_at=now,
            source_type=kwargs["source_type"],
            source_ref=kwargs["source_ref"],
        )


class Sessions:
    def __init__(self):
        self.reference = None

    async def put_dataset_reference(self, reference, *, recent_limit):
        self.reference = reference


def workbook_bytes():
    workbook = Workbook()
    first = workbook.active
    first.title = "销售"
    first.append(["销售月报"])
    first.append(["月份", "销售额"])
    first.append(["2026-07", 100])
    second = workbook.create_sheet("预算")
    second.append(["月份", "预算额"])
    second.append(["2026-07", 120])
    output = io.BytesIO()
    workbook.save(output)
    return output.getvalue()


@pytest.mark.asyncio
async def test_multi_sheet_file_is_registered_as_scoped_dataset():
    sessions = Sessions()
    importer = SpreadsheetFileImporter(
        Minio(workbook_bytes()),
        Store(),
        sessions,
        bucket="bam",
        max_bytes=1024 * 1024,
        max_rows=100,
        max_sheets=10,
        ttl_seconds=3600,
        recent_limit=10,
    )
    reference, sheets = await importer.import_object(
        object_name="uploads/report.xlsx",
        scope=DatasetScope("tenant", "user", "app", "conversation"),
    )
    assert sheets == ["销售", "预算"]
    assert reference.row_count == 2
    assert reference.columns == ("_sheet_name", "月份", "销售额", "预算额")
    assert sessions.reference["dataset_id"] == reference.dataset_id


@pytest.mark.asyncio
async def test_document_is_not_misread_as_structured_data():
    importer = SpreadsheetFileImporter(
        Minio(b"pdf"), Store(), Sessions(), bucket="bam",
        max_bytes=1024, max_rows=10, max_sheets=2, ttl_seconds=60, recent_limit=2,
    )
    with pytest.raises(FileImportError, match="knowledge-base"):
        await importer.import_object(
            object_name="uploads/manual.pdf",
            scope=DatasetScope("tenant", "user", "app", "conversation"),
        )


def test_named_sheet_is_selected_without_requiring_filter_wording():
    operation = plan_dataset_followup(
        "分析销售Sheet的趋势",
        ["_sheet_name", "月份", "销售额"],
        [
            {"_sheet_name": "销售", "月份": "7月", "销售额": 100},
            {"_sheet_name": "预算", "月份": "7月", "销售额": 120},
        ],
    )
    assert operation == {
        "type": "filter", "field": "_sheet_name", "operator": "eq", "value": "销售"
    }


def test_xlsx_archive_with_excessive_expansion_is_rejected_before_openpyxl():
    payload = io.BytesIO()
    with zipfile.ZipFile(payload, "w", compression=zipfile.ZIP_DEFLATED) as archive:
        archive.writestr("xl/worksheets/sheet1.xml", b"0" * (65 * 1024 * 1024))
    importer = SpreadsheetFileImporter(
        Minio(payload.getvalue()), Store(), Sessions(), bucket="bam",
        max_bytes=1024 * 1024, max_rows=100, max_sheets=10,
        ttl_seconds=60, recent_limit=2,
    )
    with pytest.raises(FileImportError, match="safety limit|compression ratio"):
        importer._parse_xlsx(payload.getvalue())


@pytest.mark.asyncio
async def test_xlsx_row_limit_is_enforced_across_sheets():
    workbook = Workbook()
    first = workbook.active
    first.append(["value"])
    first.append([1])
    first.append([2])
    second = workbook.create_sheet("second")
    second.append(["value"])
    second.append([3])
    output = io.BytesIO()
    workbook.save(output)
    importer = SpreadsheetFileImporter(
        Minio(output.getvalue()), Store(), Sessions(), bucket="bam",
        max_bytes=1024 * 1024, max_rows=2, max_sheets=10,
        ttl_seconds=60, recent_limit=2,
    )
    with pytest.raises(FileImportError, match="row limit"):
        await importer.import_object(
            object_name="uploads/too-many.xlsx",
            scope=DatasetScope("tenant", "user", "app", "conversation"),
        )


def test_csv_column_and_cell_text_limits_are_enforced():
    importer = SpreadsheetFileImporter(
        Minio(b"a,b,c\n1,2,3\n"), Store(), Sessions(), bucket="bam",
        max_bytes=1024, max_rows=10, max_sheets=2, ttl_seconds=60,
        recent_limit=2, max_columns=2,
    )
    with pytest.raises(FileImportError, match="column limit"):
        importer._parse_csv(b"a,b,c\n1,2,3\n")

    importer = SpreadsheetFileImporter(
        Minio(b"value\nabcdef\n"), Store(), Sessions(), bucket="bam",
        max_bytes=1024, max_rows=10, max_sheets=2, ttl_seconds=60,
        recent_limit=2, max_cell_chars=5,
    )
    with pytest.raises(FileImportError, match="text limit"):
        importer._parse_csv(b"value\nabcdef\n")


def test_named_sheet_filter_is_composed_with_aggregation():
    operation = plan_dataset_followup(
        "销售Sheet的销售额合计是多少",
        ["_sheet_name", "月份", "销售额"],
        [
            {"_sheet_name": "销售", "月份": "6月", "销售额": 80},
            {"_sheet_name": "销售", "月份": "7月", "销售额": 100},
            {"_sheet_name": "预算", "月份": "7月", "销售额": 120},
        ],
    )
    assert operation == {
        "type": "pipeline",
        "operations": [
            {"type": "filter", "field": "_sheet_name", "operator": "eq", "value": "销售"},
            {
                "type": "aggregate",
                "group_by": [],
                "aggregations": [
                    {"field": "销售额", "function": "sum", "alias": "sum_销售额"}
                ],
            },
        ],
    }


def test_sheet_name_inside_metric_does_not_silently_filter_workbook():
    operation = plan_dataset_followup(
        "分析销售额趋势",
        ["_sheet_name", "月份", "销售额"],
        [
            {"_sheet_name": "销售", "月份": "6月", "销售额": 80},
            {"_sheet_name": "预算", "月份": "7月", "销售额": 120},
        ],
    )
    assert operation is None


def test_longest_explicit_sheet_name_is_selected():
    operation = plan_dataset_followup(
        "分析销售明细Sheet",
        ["_sheet_name", "销售额"],
        [
            {"_sheet_name": "销售", "销售额": 80},
            {"_sheet_name": "销售明细", "销售额": 100},
        ],
    )
    assert operation == {
        "type": "filter", "field": "_sheet_name",
        "operator": "eq", "value": "销售明细",
    }


def test_bare_result_operations_are_detected_without_trusting_intent_label():
    assert is_dataset_operation_followup("最高的是哪个，最低的是哪个，差多少")
    assert is_dataset_operation_followup("按销售额从高到低排序")
    assert is_dataset_operation_followup("只保留华东")
    assert is_dataset_operation_followup("把刚才结果导出成Excel")
    assert is_dataset_operation_followup("展示前20条数据")
    assert not is_dataset_operation_followup("查询2026年7月销售额最高的门店")


def test_plain_preview_limit_preserves_source_order_without_numeric_sort():
    rows = [
        {"经销商名称": "甲", "销售额": 10},
        {"经销商名称": "乙", "销售额": 30},
    ]
    assert plan_dataset_followup(
        "展示前20条数据", ["经销商名称", "销售额"], rows
    ) == {"type": "limit", "count": 20}


def test_extrema_difference_keeps_corresponding_dimension_rows():
    operation = plan_dataset_followup(
        "最高的是哪个，最低的是哪个，差多少",
        ["订单状态", "销售额"],
        [
            {"订单状态": "已支付", "销售额": 120},
            {"订单状态": "已退款", "销售额": 80},
            {"订单状态": "待支付", "销售额": 100},
        ],
    )
    assert operation == {
        "type": "extrema",
        "field": "销售额",
        "include_difference": True,
        "result_field": "销售额差额",
    }


def test_single_extremum_returns_the_whole_matching_row():
    assert plan_dataset_followup(
        "最高的是哪个",
        ["门店", "销售额"],
        [{"门店": "A", "销售额": 1}, {"门店": "B", "销售额": 2}],
    ) == {
        "type": "sort_limit", "field": "销售额", "descending": True, "count": 1,
    }


def test_projection_and_directional_sort_are_local_operations():
    rows = [{"门店": "A", "销售额": 1, "订单量": 2}]
    assert plan_dataset_followup("只显示门店和销售额", ["门店", "销售额", "订单量"], rows) == {
        "type": "select", "columns": ["门店", "销售额"],
    }
    assert plan_dataset_followup("销售额从低到高", ["门店", "销售额"], rows) == {
        "type": "sort", "field": "销售额", "descending": False,
    }


def test_only_keep_metric_is_column_projection_not_row_filter():
    rows = [{"商品": "A", "含税销售总额": 100, "订单笔数": 2}]

    assert plan_dataset_followup(
        "只保留订单笔数，其他条件不变。",
        ["商品", "含税销售总额", "订单笔数"],
        rows,
    ) == {"type": "select", "columns": ["订单笔数"]}


def test_projection_with_missing_enrichment_column_returns_to_query_path():
    rows = [{"医院名称": "A医院"}, {"医院名称": "B医院"}]
    assert plan_dataset_followup(
        "只保留医院名称和医院等级", ["医院名称"], rows
    ) is None


def test_unit_price_comparison_uses_sales_divided_by_order_count():
    operation = plan_dataset_followup(
        "算客单价谁高",
        ["月份", "销售额", "订单量"],
        [
            {"月份": "6月", "销售额": 100, "订单量": 10},
            {"月份": "7月", "销售额": 150, "订单量": 10},
        ],
    )
    assert operation == {
        "type": "pipeline",
        "operations": [
            {
                "type": "derive", "left_field": "销售额", "right_field": "订单量",
                "operator": "divide", "result_field": "客单价",
            },
            {"type": "sort", "field": "客单价", "descending": True},
            {"type": "limit", "count": 1},
        ],
    }


def test_unit_price_calculation_fails_closed_when_source_columns_are_ambiguous():
    assert plan_dataset_followup(
        "算客单价谁高",
        ["销售额", "含税销售额", "订单量"],
        [{"销售额": 1, "含税销售额": 2, "订单量": 1}],
    ) is None


def test_bare_top_count_on_single_name_column_is_a_preview_limit():
    assert plan_dataset_followup(
        "前五个",
        ["经销商名称", "_source_row"],
        [
            {"经销商名称": "甲", "_source_row": 1},
            {"经销商名称": "乙", "_source_row": 2},
        ],
    ) == {"type": "limit", "count": 5}


def test_limit_replacement_preserves_existing_multi_metric_ranking_order():
    rows = [
        {"经销商": "甲", "销售额": 100, "合作次数": 3},
        {"经销商": "乙", "销售额": 90, "合作次数": 8},
    ]

    assert plan_dataset_followup(
        "改成前3名，其他条件不变。",
        ["经销商", "销售额", "合作次数"],
        rows,
    ) == {"type": "limit", "count": 3}


def test_preview_limit_ignores_unchanged_filter_scope_suffix():
    assert plan_dataset_followup(
        "只返回前5个，其他筛选条件不变。",
        ["经销商名称"],
        [{"经销商名称": "甲"}, {"经销商名称": "乙"}],
    ) == {"type": "limit", "count": 5}
