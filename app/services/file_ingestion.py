"""Import trusted MinIO spreadsheet objects as immutable conversational datasets."""
from __future__ import annotations

import csv
import asyncio
import io
import zipfile
import hashlib
from datetime import date, datetime, timezone
from pathlib import PurePosixPath
from typing import Any

from minio_followup_store import DatasetScope, HybridMinioFollowupStore

from app.stores import SessionStore


class FileImportError(ValueError):
    pass


class SpreadsheetFileImporter:
    def __init__(
        self,
        minio_client,
        dataset_store: HybridMinioFollowupStore,
        sessions: SessionStore,
        *,
        bucket: str,
        max_bytes: int,
        max_rows: int,
        max_sheets: int,
        ttl_seconds: int,
        recent_limit: int,
        max_columns: int = 500,
        max_cells: int = 5_000_000,
        max_cell_chars: int = 32_767,
    ) -> None:
        self.client = minio_client
        self.dataset_store = dataset_store
        self.sessions = sessions
        self.bucket = bucket
        self.max_bytes = max_bytes
        self.max_rows = max_rows
        self.max_sheets = max_sheets
        self.ttl_seconds = ttl_seconds
        self.recent_limit = recent_limit
        self.max_columns = max_columns
        self.max_cells = max_cells
        self.max_cell_chars = max_cell_chars

    @staticmethod
    def _safe_object_name(value: str) -> str:
        name = value.strip().replace("\\", "/")
        path = PurePosixPath(name)
        if not name or name.startswith("/") or ".." in path.parts or len(name) > 1024:
            raise FileImportError("invalid MinIO object name")
        return name

    @staticmethod
    def _object_version(stat: Any) -> tuple[int, str, str]:
        return (
            int(getattr(stat, "size", 0) or 0),
            str(getattr(stat, "etag", "") or "").strip('"'),
            str(getattr(stat, "version_id", "") or ""),
        )

    def _download(self, object_name: str) -> tuple[bytes, dict[str, str]]:
        stat = self.client.stat_object(self.bucket, object_name)
        before = self._object_version(stat)
        size = before[0]
        if size <= 0:
            raise FileImportError("uploaded file is empty")
        if size > self.max_bytes:
            raise FileImportError("uploaded file exceeds the configured size limit")
        response = self.client.get_object(self.bucket, object_name)
        try:
            payload = response.read(self.max_bytes + 1)
        finally:
            response.close()
            response.release_conn()
        if len(payload) > self.max_bytes:
            raise FileImportError("uploaded file exceeds the configured size limit")
        after = self._object_version(self.client.stat_object(self.bucket, object_name))
        if before != after:
            raise FileImportError("uploaded file changed while it was being imported; retry")
        return payload, {
            "etag": before[1],
            "version_id": before[2],
            "content_sha256": hashlib.sha256(payload).hexdigest(),
        }

    @staticmethod
    def _cell(value: Any) -> Any:
        if isinstance(value, (date, datetime)):
            return value.isoformat()
        return value

    @staticmethod
    def _headers(values: list[Any]) -> list[str]:
        output: list[str] = []
        counts: dict[str, int] = {}
        for index, value in enumerate(values, 1):
            base = str(value).strip() if value is not None else ""
            base = base or f"column_{index}"
            counts[base] = counts.get(base, 0) + 1
            output.append(base if counts[base] == 1 else f"{base}_{counts[base]}")
        return output

    def _validate_row_shape(self, values: list[Any]) -> int:
        if len(values) > self.max_columns:
            raise FileImportError("spreadsheet exceeds the configured column limit")
        for value in values:
            if isinstance(value, str) and len(value) > self.max_cell_chars:
                raise FileImportError("spreadsheet cell exceeds the configured text limit")
        return len(values)

    @staticmethod
    def _detect_header_index(rows: list[list[Any]]) -> int:
        """Prefer a dense, textual header over title/blank preamble rows."""
        if not rows:
            raise FileImportError("spreadsheet contains no rows")
        candidates = []
        for index, values in enumerate(rows[:20]):
            non_empty = [value for value in values if value is not None and str(value).strip()]
            textual = sum(isinstance(value, str) for value in non_empty)
            unique = len({str(value).strip() for value in non_empty})
            # Earlier rows win only when their structural score is equal.
            candidates.append(((len(non_empty), textual, unique, -index), index))
        return max(candidates)[1]

    def _parse_csv(self, payload: bytes) -> tuple[list[str], list[dict[str, Any]], list[str]]:
        text = None
        for encoding in ("utf-8-sig", "gb18030"):
            try:
                text = payload.decode(encoding)
                break
            except UnicodeDecodeError:
                continue
        if text is None:
            raise FileImportError("CSV encoding must be UTF-8 or GB18030")
        sample = text[:8192]
        try:
            dialect = csv.Sniffer().sniff(sample, delimiters=",\t;|")
        except csv.Error:
            dialect = csv.excel
        reader = csv.reader(io.StringIO(text), dialect)
        preamble: list[list[Any]] = []
        cell_count = 0
        for _ in range(20):
            try:
                values = list(next(reader))
                cell_count += self._validate_row_shape(values)
                if cell_count > self.max_cells:
                    raise FileImportError("spreadsheet exceeds the configured cell limit")
                preamble.append(values)
            except StopIteration:
                break
        if not preamble:
            raise FileImportError("CSV contains no header row")
        header_index = self._detect_header_index(preamble)
        columns = self._headers(preamble[header_index])
        rows = []
        for values in preamble[header_index + 1 :]:
            if not any(str(value).strip() for value in values):
                continue
            values = list(values[: len(columns)]) + [None] * max(0, len(columns) - len(values))
            rows.append(dict(zip(columns, values)))
            if len(rows) > self.max_rows:
                raise FileImportError("spreadsheet exceeds the configured row limit")
        for values in reader:
            values = list(values)
            cell_count += self._validate_row_shape(values)
            if cell_count > self.max_cells:
                raise FileImportError("spreadsheet exceeds the configured cell limit")
            if not any(str(value).strip() for value in values):
                continue
            values = list(values[: len(columns)]) + [None] * max(0, len(columns) - len(values))
            rows.append(dict(zip(columns, values)))
            if len(rows) > self.max_rows:
                raise FileImportError("spreadsheet exceeds the configured row limit")
        return columns, rows, ["CSV"]

    def _validate_xlsx_archive(self, payload: bytes) -> None:
        """Reject malformed/encrypted XLSX archives and bounded zip bombs."""
        try:
            with zipfile.ZipFile(io.BytesIO(payload)) as archive:
                infos = archive.infolist()
                if not infos or any(info.flag_bits & 0x1 for info in infos):
                    raise FileImportError("encrypted or empty XLSX files are not supported")
                expanded = sum(max(0, info.file_size) for info in infos)
                expanded_limit = max(64 * 1024 * 1024, min(self.max_bytes * 5, 512 * 1024 * 1024))
                if expanded > expanded_limit:
                    raise FileImportError("XLSX expanded content exceeds the safety limit")
                compressed = sum(max(1, info.compress_size) for info in infos)
                if expanded > 64 * 1024 * 1024 and expanded / compressed > 200:
                    raise FileImportError("XLSX compression ratio exceeds the safety limit")
        except FileImportError:
            raise
        except zipfile.BadZipFile as exc:
            raise FileImportError("invalid or unsupported XLSX file") from exc

    def _parse_xlsx(self, payload: bytes) -> tuple[list[str], list[dict[str, Any]], list[str]]:
        from openpyxl import load_workbook

        self._validate_xlsx_archive(payload)
        try:
            workbook = load_workbook(io.BytesIO(payload), read_only=True, data_only=True)
        except Exception as exc:
            raise FileImportError("invalid or unsupported XLSX file") from exc
        if len(workbook.sheetnames) > self.max_sheets:
            workbook.close()
            raise FileImportError("workbook exceeds the configured sheet limit")
        rows: list[dict[str, Any]] = []
        all_columns: list[str] = ["_sheet_name"]
        sheet_rows: list[tuple[str, list[str], list[list[Any]]]] = []
        total_rows = 0
        estimated_cells = 0
        try:
            for sheet in workbook.worksheets:
                if sheet.max_column > self.max_columns:
                    raise FileImportError("workbook exceeds the configured column limit")
                if sheet.max_row > self.max_rows + 20:
                    raise FileImportError("workbook exceeds the configured row limit")
                estimated_cells += sheet.max_row * max(1, sheet.max_column)
                if estimated_cells > self.max_cells:
                    raise FileImportError("workbook exceeds the configured cell limit")
                iterator = (
                    [self._cell(value) for value in values]
                    for values in sheet.iter_rows(values_only=True)
                    if any(value is not None for value in values)
                )
                preamble: list[list[Any]] = []
                for _ in range(20):
                    try:
                        preamble.append(next(iterator))
                    except StopIteration:
                        break
                if not preamble:
                    continue
                for values in preamble:
                    self._validate_row_shape(values)
                header_index = self._detect_header_index(preamble)
                columns = self._headers(preamble[header_index])
                for column in columns:
                    if column not in all_columns:
                        all_columns.append(column)
                values_list = list(preamble[header_index + 1 :])
                for values in iterator:
                    self._validate_row_shape(values)
                    values_list.append(values)
                    if total_rows + len(values_list) > self.max_rows:
                        raise FileImportError("workbook exceeds the configured row limit")
                sheet_rows.append((sheet.title, columns, values_list))
                total_rows += len(values_list)
                if total_rows > self.max_rows:
                    raise FileImportError("workbook exceeds the configured row limit")
        finally:
            workbook.close()
        for sheet_name, columns, values_list in sheet_rows:
            for values in values_list:
                values = values[: len(columns)] + [None] * max(0, len(columns) - len(values))
                row = {column: None for column in all_columns}
                row["_sheet_name"] = sheet_name
                row.update(dict(zip(columns, values)))
                rows.append(row)
        return all_columns, rows, [item[0] for item in sheet_rows]

    async def import_object(
        self,
        *,
        object_name: str,
        scope: DatasetScope,
        source_type: str = "UPLOADED_SPREADSHEET",
        semantic_model_id: int | None = None,
        business_domain_ids: list[int] | None = None,
    ):
        object_name = self._safe_object_name(object_name)
        suffix = PurePosixPath(object_name).suffix.lower()
        if suffix not in {".csv", ".xlsx"}:
            raise FileImportError(
                "only CSV/XLSX can be imported as structured datasets; documents must use knowledge-base ingestion"
            )
        payload, source_identity = await asyncio.to_thread(self._download, object_name)
        columns, rows, sheets = await asyncio.to_thread(
            self._parse_csv if suffix == ".csv" else self._parse_xlsx,
            payload,
        )
        if not rows:
            raise FileImportError("spreadsheet contains no data rows")
        now = datetime.now(timezone.utc)
        reference = await asyncio.to_thread(
            self.dataset_store.save_dataset,
            scope=scope,
            columns=columns,
            rows=rows,
            snapshot_id=f"upload-sha256:{source_identity['content_sha256']}",
            data_as_of=now,
            source_type=source_type,
            source_ref=(
                f"minio://{self.bucket}/{object_name}"
                f"#sha256={source_identity['content_sha256']}"
            ),
            ttl_seconds=self.ttl_seconds,
            semantic_model_id=semantic_model_id,
            business_domain_ids=business_domain_ids or [],
            transformation_log=({
                "type": "FILE_IMPORT",
                "sheets": sheets,
                "source_etag": source_identity["etag"],
                "source_version_id": source_identity["version_id"],
                "content_sha256": source_identity["content_sha256"],
            },),
        )
        await self.sessions.put_dataset_reference(
            reference.to_dict(), recent_limit=self.recent_limit
        )
        return reference, sheets
